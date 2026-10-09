"""Page-tagged text layer for stored documents (M1 spec §3.1)."""

import io
import re
import unicodedata
from dataclasses import dataclass
from typing import Literal

MIN_PAGE_CHARS = 30
Kind = Literal["pdf", "html", "docx", "text", "other"]
_GENERIC_MIME = {"", "application/octet-stream", "binary/octet-stream"}


@dataclass(frozen=True)
class Page:
    number: int  # 1-based page (PDF) or section index (HTML/DOCX/text)
    text: str
    heading: str = ""
    has_text: bool = True


def normalize(text: str | None) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text or "")).strip()


def _has_text(text: str) -> bool:
    return len(re.sub(r"\s", "", text)) >= MIN_PAGE_CHARS


def document_kind(mime: str | None, file_name: str) -> Kind:
    m = (mime or "").split(";")[0].strip().lower()
    name = (file_name or "").lower()
    if m in _GENERIC_MIME:
        m = ""
    if m == "application/pdf" or (not m and name.endswith(".pdf")):
        return "pdf"
    if m == "text/html" or (not m and name.endswith((".html", ".htm"))):
        return "html"
    if (
        m == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or (not m and name.endswith(".docx"))
    ):
        return "docx"
    if m == "text/plain" or (not m and name.endswith(".txt")):
        return "text"
    return "other"


def pdf_pages(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    pages: list[Page] = []
    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        try:
            text = normalize(page.extract_text())
        except Exception:  # a broken page must not lose the rest of the document
            text = ""
        pages.append(Page(i, text, has_text=_has_text(text)))
    return pages


def _split_markdown(md: str) -> list[Page]:
    sections: list[tuple[str, list[str]]] = [("", [])]
    for line in md.splitlines():
        m = re.match(r"^#{1,2}\s+(.*)", line)
        if m:
            sections.append((m.group(1).strip(), []))
        else:
            sections[-1][1].append(line)
    pages: list[Page] = []
    for heading, body in sections:
        text = normalize(" ".join([heading, *body]))
        if not text:
            continue
        pages.append(
            Page(len(pages) + 1, text, heading=heading, has_text=_has_text(text))
        )
    return pages or [Page(1, "", has_text=False)]


def html_sections(data: bytes) -> list[Page]:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "noscript"]):
        tag.decompose()
    return _split_markdown(markdownify(str(soup), heading_style="ATX"))


def docx_sections(data: bytes) -> list[Page]:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "").lower()
        if style in ("title", "heading 1", "heading 2"):
            lines.append(f"## {para.text}")
        else:
            lines.append(para.text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return _split_markdown("\n".join(lines))


def extract_pages(data: bytes, mime: str | None, file_name: str) -> list[Page] | None:
    """Pages for a stored document, or None when the format is unsupported."""
    kind = document_kind(mime, file_name)
    if kind == "pdf":
        return pdf_pages(data)
    if kind == "html":
        return html_sections(data)
    if kind == "docx":
        return docx_sections(data)
    if kind == "text":
        return _split_markdown(data.decode("utf-8", errors="replace"))
    return None
