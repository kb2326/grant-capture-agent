import io

import docx
from fpdf import FPDF

from ingest.textlayer import (
    MIN_PAGE_CHARS,
    document_kind,
    docx_sections,
    extract_pages,
    html_sections,
    normalize,
    pdf_pages,
)


def _pdf(pages: list[str]) -> bytes:
    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        if text:
            pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


def test_normalize_collapses_whitespace_and_nfkc():
    assert normalize("  ﬁrst\n\n  line\tend ") == "first line end"


def test_pdf_pages_are_numbered_and_blank_pages_flagged():
    data = _pdf(["Eligibility: only small business concerns may apply. " * 3, ""])
    pages = pdf_pages(data)
    assert [p.number for p in pages] == [1, 2]
    assert pages[0].has_text and "small business concerns" in pages[0].text
    assert not pages[1].has_text


def test_html_sections_split_on_headings_and_drop_scripts():
    html = (
        b"<html><script>evil()</script><h1>Overview</h1><p>Funding for grid research projects here.</p>"
        b"<h2>Eligibility</h2><p>Only nonprofit organizations are eligible to apply for this award.</p></html>"
    )
    pages = html_sections(html)
    assert [p.heading for p in pages] == ["Overview", "Eligibility"]
    assert "evil" not in " ".join(p.text for p in pages)
    assert "nonprofit organizations" in pages[1].text


def test_docx_sections_use_heading_styles_and_tables():
    d = docx.Document()
    d.add_heading("Eligibility", level=1)
    d.add_paragraph(
        "Applicants must be registered in SAM.gov before submission of the proposal."
    )
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Page limit"
    t.rows[0].cells[1].text = "15 pages"
    buf = io.BytesIO()
    d.save(buf)
    pages = docx_sections(buf.getvalue())
    assert pages[0].heading == "Eligibility"
    assert "registered in SAM.gov" in pages[0].text and "15 pages" in pages[-1].text


def test_document_kind_and_unsupported_doc():
    assert document_kind("application/pdf", "x") == "pdf"
    assert document_kind("text/html;charset=UTF-8", "x") == "html"
    assert document_kind(None, "Notice.DOCX") == "docx"
    assert document_kind("application/octet-stream", "nofo.pdf") == "pdf"
    assert extract_pages(b"\xd0\xcf\x11\xe0", "application/msword", "old.doc") is None


def test_min_page_chars_constant():
    assert MIN_PAGE_CHARS == 30
